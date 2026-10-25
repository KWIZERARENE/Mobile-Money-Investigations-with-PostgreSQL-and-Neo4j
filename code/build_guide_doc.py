import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_formatted_text(p, text, font_name="Calibri", font_size=10.5, color_rgb=(0x22, 0x22, 0x22)):
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for tok in tokens:
        if not tok:
            continue
        run = p.add_run()
        run.font.name = font_name
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor(*color_rgb)
        
        if tok.startswith('**') and tok.endswith('**') and len(tok) >= 4:
            run.text = tok[2:-2]
            run.font.bold = True
        elif tok.startswith('*') and tok.endswith('*') and len(tok) >= 2:
            run.text = tok[1:-1]
            run.font.italic = True
        elif tok.startswith('`') and tok.endswith('`') and len(tok) >= 2:
            run.text = tok[1:-1]
            run.font.name = "Consolas"
            run.font.size = Pt(font_size * 0.92)
            run.font.color.rgb = RGBColor(0x90, 0x20, 0x20)
        else:
            run.text = tok

def format_row(row, is_header=False, is_even=False):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    if is_header:
        trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

    for cell in row.cells:
        if is_header:
            set_cell_background(cell, "1F4E78")
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(9.5)
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        else:
            bg = "F4F6F9" if is_even else "FFFFFF"
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=140, right=140)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(9.0)
                    r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

def build_guide_docx(md_path, docx_path, logo_path):
    doc = docx.Document()
    for sec in doc.sections:
        sec.top_margin = Inches(0.85)
        sec.bottom_margin = Inches(0.85)
        sec.left_margin = Inches(0.9)
        sec.right_margin = Inches(0.9)

    # Logo
    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(10)
        p_logo.paragraph_format.space_after = Pt(6)
        p_logo.add_run().add_picture(logo_path, width=Inches(1.5))

    with open(md_path, "r", encoding="utf-8") as f:
        md_text = f.read()

    lines = md_text.splitlines()
    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []
    
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        # Code block
        if stripped.startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_lines = []
            else:
                in_code_block = False
                if code_lines:
                    tbl_code = doc.add_table(rows=1, cols=1)
                    tbl_code.alignment = WD_TABLE_ALIGNMENT.CENTER
                    cell = tbl_code.rows[0].cells[0]
                    cell.width = Inches(6.5)
                    set_cell_background(cell, "F8F9FA")
                    set_cell_margins(cell, top=80, bottom=80, left=140, right=140)
                    
                    tcPr = cell._tc.get_or_add_tcPr()
                    bdr = parse_xml(
                        f'<w:tcBorders {nsdecls("w")}>'
                        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>'
                        f'<w:left w:val="single" w:sz="18" w:space="0" w:color="1F4E78"/>'
                        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>'
                        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>'
                        f'</w:tcBorders>'
                    )
                    tcPr.append(bdr)
                    
                    p_c = cell.paragraphs[0]
                    p_c.paragraph_format.space_before = Pt(2)
                    p_c.paragraph_format.space_after = Pt(2)
                    p_c.paragraph_format.line_spacing = 1.05
                    
                    full_code_text = "\n".join(code_lines)
                    r_code = p_c.add_run(full_code_text)
                    r_code.font.name = "Consolas"
                    r_code.font.size = Pt(8.5)
                    r_code.font.color.rgb = RGBColor(0x24, 0x29, 0x2E)
                    
                    p_after = doc.add_paragraph()
                    p_after.paragraph_format.space_after = Pt(4)
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Table
        if stripped.startswith("|") and stripped.endswith("|"):
            if not in_table:
                in_table = True
                table_lines = [stripped]
            else:
                table_lines.append(stripped)
            i += 1
            continue
        else:
            if in_table:
                in_table = False
                rows_data = []
                for tl in table_lines:
                    cells = [c.strip() for c in tl.strip("|").split("|")]
                    if all(set(c).issubset({'-', ':', ' '}) for c in cells):
                        continue
                    rows_data.append(cells)
                    
                if rows_data:
                    n_rows = len(rows_data)
                    n_cols = max(len(r) for r in rows_data)
                    norm_rows = []
                    for r in rows_data:
                        if len(r) < n_cols:
                            r = r + [""] * (n_cols - len(r))
                        norm_rows.append(r[:n_cols])
                        
                    t_doc = doc.add_table(rows=n_rows, cols=n_cols)
                    t_doc.alignment = WD_TABLE_ALIGNMENT.CENTER
                    set_table_borders(t_doc, color="D3D3D3", sz="4", val="single")
                    
                    for r_idx, row_vals in enumerate(norm_rows):
                        table_row = t_doc.rows[r_idx]
                        is_hdr = (r_idx == 0)
                        for c_idx, val in enumerate(row_vals):
                            cell = table_row.cells[c_idx]
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.space_after = Pt(2)
                            p.paragraph_format.line_spacing = 1.05
                            add_formatted_text(p, val, font_size=9.0 if not is_hdr else 9.5)
                        format_row(table_row, is_header=is_hdr, is_even=(r_idx % 2 == 0))
                        
                    p_spacer = doc.add_paragraph()
                    p_spacer.paragraph_format.space_after = Pt(4)

        # Headings
        if stripped.startswith("# "):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(5)
            h.paragraph_format.keep_with_next = True
            r = h.add_run(stripped[2:])
            r.font.name = "Calibri"
            r.font.size = Pt(15)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
            i += 1
            continue

        if stripped.startswith("## "):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
            r = h.add_run(stripped[3:])
            r.font.name = "Calibri"
            r.font.size = Pt(12.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
            i += 1
            continue

        if stripped.startswith("### "):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(9)
            h.paragraph_format.space_after = Pt(3)
            h.paragraph_format.keep_with_next = True
            r = h.add_run(stripped[4:])
            r.font.name = "Calibri"
            r.font.size = Pt(11)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
            i += 1
            continue

        if stripped == "---":
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3.5)
        p.paragraph_format.line_spacing = 1.15
        
        if stripped.startswith("- ") or stripped.startswith("* "):
            p.paragraph_format.left_indent = Inches(0.25)
            r_bullet = p.add_run("• ")
            r_bullet.font.name = "Calibri"
            r_bullet.font.bold = True
            r_bullet.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
            add_formatted_text(p, stripped[2:])
        elif re.match(r"^\d+\.\s+", stripped):
            m = re.match(r"^(\d+\.)\s+(.*)", stripped)
            p.paragraph_format.left_indent = Inches(0.25)
            r_num = p.add_run(f"{m.group(1)} ")
            r_num.font.name = "Calibri"
            r_num.font.bold = True
            r_num.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
            add_formatted_text(p, m.group(2))
        else:
            add_formatted_text(p, stripped)

        i += 1

    doc.save(docx_path)
    print(f"Saved guide docx to: {docx_path}")

if __name__ == "__main__":
    md_file = r"c:\assignement\report\Project_Code_Explanation_Guide.md"
    docx_file = r"c:\assignement\report\Project_Code_Explanation_Guide.docx"
    logo = r"c:\assignement\report\auca_logo.png"
    build_guide_docx(md_file, docx_file, logo)
